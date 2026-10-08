"""
Budowa skoroszytu ResInvest_Magazyn_Zabrze 4.0 (XLSM) z pliku 3.x.

Zmiany wykonywane chirurgicznie w XML (bez openpyxl przy zapisie), aby zachować
model danych Power Pivot, tabele przestawne, fragmentatory i formatowanie:

  1. Tabela_dane (arkusz Dane): 8 nowych kolumn AM..AT
     ID_operacji, Transport, Trasa, Pojazd, Operator_kierowca, Rola_wiersza, Receptura, Wersja
     + migracja wszystkich wierszy historycznych (ID H-xxxx, Transport TAK/NIE, rola wiersza).
  2. Nowe arkusze: REJESTR_OPERACJI (z przyciskami), RECEPTURY, HISTORIA_ZMIAN, ARCHIWUM, DOKUMENT.
  3. MAGAZYN: statystyki transportu liczone z kolumny Transport (nie z typu "TRANSPORT").
  4. SŁOWNIK: "TRANSPORT" usunięty z listy typów operacji.
  5. Projekt VBA: nowe/zmienione moduły (src/*), projektanci formularzy bez zmian.

Użycie: python3 build_workbook.py <oryginał.xlsm> <katalog src> <wynik.xlsm>
"""
import sys, os, re, zipfile, uuid, datetime, html
from collections import defaultdict, OrderedDict
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vbaproj

NOWE_KOLUMNY = ["ID_operacji", "Transport", "Trasa", "Pojazd", "Operator_kierowca", "Rola_wiersza", "Receptura", "Wersja"]
KOL_LITERY = ["AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT"]
GLOWNY, Z_AUTO, P_AUTO, S_AUTO = "GŁÓWNY", "ZUŻYCIE AUTO", "PRODUKCJA AUTO", "SPRZEDAŻ AUTO"

RECEPTURY = [
    # produkt wyjściowy, JM produktu, surowiec, JM surowca, przelicznik (produktu z 1 jedn. surowca), aktywna, uwagi
    ("Zrębka Produkcyjna Leśna", "MP", "Drewno opałowe z lasu", "M3", 4, "TAK", "1 m³ drewna = 4 MP zrębki"),
    ("Zrębka Produkcyjna Inwestycyjna", "MP", "Drewno przemysłowe z wycinek inwest.", "M3", 4, "TAK", "1 m³ drewna = 4 MP zrębki"),
    ("Zrębka Produkcyjna Drzewna", "MP", "Drewno opałowe z lasu", "M3", 4, "TAK", "sprawdź surowiec dla tej zrębki"),
]

ST_NAGLOWEK, ST_TEKST, ST_LICZBA, ST_DATA, ST_CZAS = 17, 20, 23, 32, 26


def esc(s):
    return html.escape(str(s), quote=False)


def col_num(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def col_letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def excel_serial(d):
    if isinstance(d, datetime.datetime):
        delta = d - datetime.datetime(1899, 12, 30)
        return delta.days + delta.seconds / 86400
    if isinstance(d, datetime.date):
        return (d - datetime.date(1899, 12, 30)).days
    return None


def cell(ref, value, style):
    if value is None or value == "":
        return '<c r="%s" s="%d"/>' % (ref, style)
    if isinstance(value, (int, float)):
        return '<c r="%s" s="%d"><v>%s</v></c>' % (ref, style, repr(value) if isinstance(value, float) else value)
    return '<c r="%s" s="%d" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, style, esc(value))


def worksheet(code_name, cols, rows_xml, extra_after_sheetdata="", tab_color=None, drawing_rid=None, freeze=None):
    tab = '<tabColor rgb="%s"/>' % tab_color if tab_color else ""
    pane = ""
    if freeze:
        pane = '<pane ySplit="%d" topLeftCell="A%d" activePane="bottomLeft" state="frozen"/>' % (freeze - 1, freeze)
    colxml = "".join('<col min="%d" max="%d" width="%s" customWidth="1"/>' % (i, i, w) for i, w in cols)
    drw = '<drawing r:id="%s"/>' % drawing_rid if drawing_rid else ""
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheetPr codeName="%s">%s</sheetPr><sheetViews><sheetView workbookViewId="0">%s</sheetView></sheetViews>'
            '<sheetFormatPr defaultRowHeight="15"/><cols>%s</cols><sheetData>%s</sheetData>%s'
            '<pageMargins left="0.5" right="0.5" top="0.6" bottom="0.6" header="0.3" footer="0.3"/>'
            '<pageSetup paperSize="9" orientation="portrait" fitToHeight="0"/>%s</worksheet>'
            % (code_name, tab, pane, colxml, rows_xml, extra_after_sheetdata, drw))


def row_xml(r, cells, ht=None):
    h = ' ht="%s" customHeight="1"' % ht if ht else ""
    return '<row r="%d"%s>%s</row>' % (r, h, "".join(cells))


# ---------------------------------------------------------------------------------------------------------
def migracja(dane_rows):
    """dane_rows: lista krotek wartości wierszy 2..N (kolumny A..AL). Zwraca listę słowników nowych kolumn."""
    out = [None] * len(dane_rows)
    auto = lambda r: "utomatyczn" in str(r[22] or "")
    grupy = defaultdict(list)
    for i, r in enumerate(dane_rows):
        dod = r[24]
        klucz = (r[5], r[2], dod.replace(second=0, microsecond=0) if isinstance(dod, datetime.datetime) else dod)
        grupy[klucz].append(i)
    nr = 0
    przypisane = set()
    # operacje złożone: dokładnie jeden wiersz główny + wiersze "Automatyczn..." o tym samym nr WZ, dacie i minucie wpisu
    for idx in sorted(grupy.values(), key=lambda v: min(v)):
        glowne = [i for i in idx if not auto(dane_rows[i])]
        auta = [i for i in idx if auto(dane_rows[i])]
        if len(idx) > 1 and len(glowne) == 1 and auta:
            nr += 1
            oid = "H-%04d" % nr
            for i in idx:
                przypisane.add(i)
                r = dane_rows[i]
                rola = GLOWNY if i == glowne[0] else {"ZUŻYCIE": Z_AUTO, "PRODUKCJA": P_AUTO, "SPRZEDAŻ": S_AUTO}.get(str(r[4]), GLOWNY)
                out[i] = (oid, rola)
    for i, r in enumerate(dane_rows):
        if i in przypisane:
            continue
        nr += 1
        out[i] = ("H-%04d" % nr, GLOWNY)
    # numeracja wg kolejności wierszy (czytelniejsza): przemapuj ID wg pierwszego wystąpienia
    mapa, k = {}, 0
    for i in range(len(dane_rows)):
        oid = out[i][0]
        if oid not in mapa:
            k += 1
            mapa[oid] = "H-%04d" % k
    wynik = []
    for i, r in enumerate(dane_rows):
        oid, rola = mapa[out[i][0]], out[i][1]
        ma_tr = any(v not in (None, "", 0) for v in (r[17], r[18], r[19], r[20]))
        wynik.append({"ID_operacji": oid, "Transport": "TAK" if ma_tr else "NIE", "Trasa": "", "Pojazd": "",
                      "Operator_kierowca": "", "Rola_wiersza": rola, "Receptura": "", "Wersja": 1})
    # transport operacji złożonej = transport wiersza głównego
    tr_op = {w["ID_operacji"]: w["Transport"] for w in wynik if w["Rola_wiersza"] == GLOWNY}
    for w in wynik:
        w["Transport"] = tr_op.get(w["ID_operacji"], w["Transport"])
    return wynik


def rejestr_wiersze(dane_rows, migr):
    """Wiersze arkusza REJESTR_OPERACJI (jak modMagazyn.OdswiezRejestr)."""
    by_id = defaultdict(list)
    for r, m in zip(dane_rows, migr):
        by_id[m["ID_operacji"]].append((r, m))
    out = []
    for r, m in reversed(list(zip(dane_rows, migr))):
        if m["Rola_wiersza"] != GLOWNY:
            continue
        typ = str(r[4] or "")
        sur = zuz = jm = None
        for r2, m2 in by_id[m["ID_operacji"]]:
            if r2 is r:
                continue
            t2 = str(r2[4] or "")
            if t2 == "ZUŻYCIE":
                sur, zuz, jm = r2[13], r2[8], r2[9]
            elif t2 == "PRODUKCJA":
                typ += " + PRODUKCJA"
            elif t2 == "SPRZEDAŻ":
                typ += " + SPRZEDAŻ"
        przew = " / ".join(x for x in [str(r[17] or ""), str(r[18] or "")] if x)
        out.append([r[2], typ, r[13], r[8], r[9], "Tak" if m["Transport"] == "TAK" else "Nie", przew, r[20],
                    sur, zuz, jm, m["ID_operacji"], 1, r[22]])
    return out


# ---------------------------------------------------------------------------------------------------------
def build(orig, src_dir, out_path):
    zin = zipfile.ZipFile(orig)
    files = OrderedDict((i.filename, zin.read(i.filename)) for i in zin.infolist())
    infos = {i.filename: i for i in zin.infolist()}

    wb = openpyxl.load_workbook(orig, data_only=True)   # tylko odczyt wartości do migracji
    ws = wb["Dane"]
    dane_rows = [tuple(c for c in row[:38]) for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True)]
    ostatni = ws.max_row
    migr = migracja(dane_rows)

    # ---------------- 1. arkusz Dane: nowe kolumny + migracja ----------------
    s = files["xl/worksheets/sheet1.xml"].decode("utf-8")

    def dodaj_do_wiersza(m):
        rx = m.group(0)
        r = int(re.search(r'<row r="(\d+)"', rx).group(1))
        if r == 1:
            nowe = "".join(cell("%s1" % L, n, ST_NAGLOWEK) for L, n in zip(KOL_LITERY, NOWE_KOLUMNY))
        elif 2 <= r <= ostatni:
            w = migr[r - 2]
            nowe = "".join(cell("%s%d" % (L, r), w[n], ST_LICZBA if n == "Wersja" else ST_TEKST) for L, n in zip(KOL_LITERY, NOWE_KOLUMNY))
        else:
            return rx
        rx = re.sub(r'spans="\d+:\d+"', 'spans="1:46"', rx, count=1)
        return rx.replace("</row>", nowe + "</row>")
    s, n = re.subn(r'<row r="\d+"[^>]*>.*?</row>', dodaj_do_wiersza, s, flags=re.S)
    assert n >= ostatni, (n, ostatni)
    s = s.replace('<dimension ref="A1:AL%d"/>' % ostatni, '<dimension ref="A1:AT%d"/>' % ostatni)
    s = s.replace("</cols>", '<col min="39" max="39" width="13" style="22" customWidth="1"/><col min="40" max="40" width="11" style="22" customWidth="1"/>'
                             '<col min="41" max="43" width="20" style="22" customWidth="1"/><col min="44" max="44" width="16" style="22" customWidth="1"/>'
                             '<col min="45" max="45" width="40" style="22" customWidth="1"/><col min="46" max="46" width="8" style="22" customWidth="1"/></cols>', 1)
    files["xl/worksheets/sheet1.xml"] = s.encode("utf-8")

    t = files["xl/tables/table1.xml"].decode("utf-8")
    t = t.replace('ref="A1:AL%d"' % ostatni, 'ref="A1:AT%d"' % ostatni).replace('ref="A2:AL%d"' % ostatni, 'ref="A2:AT%d"' % ostatni)
    ids = [int(x) for x in re.findall(r'<tableColumn id="(\d+)"', t)]
    nid = max(ids)
    nowe = ""
    for nazwa in NOWE_KOLUMNY:
        nid += 1
        nowe += '<tableColumn id="%d" xr3:uid="{%s}" name="%s" dataCellStyle="Normal"/>' % (nid, str(uuid.uuid4()).upper(), nazwa)
    t = t.replace('<tableColumns count="38">', '<tableColumns count="46">').replace("</tableColumns>", nowe + "</tableColumns>")
    files["xl/tables/table1.xml"] = t.encode("utf-8")

    # ---------------- 2. MAGAZYN: transport z kolumny Transport ----------------
    s2 = files["xl/worksheets/sheet2.xml"].decode("utf-8")
    licz = sum(1 for m in migr if m["Transport"] == "TAK" and m["Rola_wiersza"] == GLOWNY)
    km = sum((r[19] or 0) for r, m in zip(dane_rows, migr) if m["Transport"] == "TAK" and isinstance(r[19], (int, float)))
    koszt = sum((r[20] or 0) for r, m in zip(dane_rows, migr) if m["Transport"] == "TAK" and isinstance(r[20], (int, float)))
    zam = {
        "AA12": ('COUNTIFS(Dane!$AN:$AN,"TAK",Dane!$AR:$AR,"GŁÓWNY")', licz),
        "AB12": ('SUMIFS(Dane!$T:$T,Dane!$AN:$AN,"TAK")', km),
        "AC12": ('SUMIFS(Dane!$U:$U,Dane!$AN:$AN,"TAK")', koszt),
    }
    for ref, (f, v) in zam.items():
        s2, k = re.subn(r'<c r="%s" s="20"><f>[^<]*</f><v/></c>' % ref, '<c r="%s" s="20"><f>%s</f><v>%s</v></c>' % (ref, esc(f), v), s2)
        assert k == 1, ref
    s2 = s2.replace("<t>Transport nie wpływa na stan magazynu</t>", "<t>Transport = pole w operacji (TAK/NIE); nie wpływa na stan magazynu</t>")
    files["xl/worksheets/sheet2.xml"] = s2.encode("utf-8")

    # ---------------- 3. SŁOWNIK: bez typu TRANSPORT ----------------
    s5 = files["xl/worksheets/sheet5.xml"].decode("utf-8")
    s5, k = re.subn(r'<c r="F7" t="inlineStr"><is><t>TRANSPORT</t></is></c>', "", s5)
    assert k == 1
    files["xl/worksheets/sheet5.xml"] = s5.encode("utf-8")

    # ---------------- 4. nowe arkusze ----------------
    teraz = datetime.datetime.now().replace(microsecond=0)
    nowe_ark = []   # (nazwa, plik, codeName, xml, rels)

    # REJESTR_OPERACJI
    rej = rejestr_wiersze(dane_rows, migr)
    nag = ["Data", "Typ", "Materiał", "Ilość", "Jednostka", "Transport", "Przewoźnik / nr rej.", "Koszt transportu",
           "Surowiec zużyty", "Zużycie", "JM surowca", "ID operacji", "Wersja", "Uwagi"]
    rows = [row_xml(1, [cell("A1", "REJESTR OPERACJI — jedna linia = jedna operacja (transport jest częścią operacji)", ST_NAGLOWEK)], 26),
            row_xml(2, [cell("A2", "Zaznacz wiersz operacji i użyj przycisku: KOREKTA / USUŃ / DOKUMENT PDF. Lista odświeża się po każdej zmianie i przy wejściu na arkusz.", ST_TEKST)]),
            row_xml(3, [cell("B3", "Operacji: %d   |   odświeżono przy migracji %s" % (len(rej), teraz.strftime("%d.%m.%Y %H:%M")), ST_TEKST)]),
            row_xml(5, [cell("%s5" % col_letters(i + 1), h, ST_NAGLOWEK) for i, h in enumerate(nag)], 30)]
    for k, w in enumerate(rej):
        r = 6 + k
        cs = []
        for j, v in enumerate(w):
            ref = "%s%d" % (col_letters(j + 1), r)
            if j == 0:
                cs.append(cell(ref, excel_serial(v), ST_DATA))
            elif j in (3, 7, 9) and isinstance(v, (int, float)):
                cs.append(cell(ref, float(v), ST_LICZBA))
            else:
                cs.append(cell(ref, "" if v is None else v, ST_TEKST))
        rows.append(row_xml(r, cs))
    nowe_ark.append(("REJESTR_OPERACJI", "sheet6.xml", "Arkusz6",
                     worksheet("Arkusz6", [(1, 12), (2, 22), (3, 30), (4, 10), (5, 10), (6, 10), (7, 28), (8, 14), (9, 30), (10, 10), (11, 10), (12, 12), (13, 8), (14, 40)],
                               "".join(rows), tab_color="FF1E6B45", drawing_rid="rId1", freeze=6), True))

    # RECEPTURY
    nag = ["Produkt wyjściowy", "JM produktu", "Surowiec (zużywany)", "JM surowca", "Przelicznik: ilość produktu z 1 jedn. surowca", "Aktywna (TAK/NIE)", "Uwagi"]
    rows = [row_xml(1, [cell("%s1" % col_letters(i + 1), h, ST_NAGLOWEK) for i, h in enumerate(nag)], 30)]
    for k, rec in enumerate(RECEPTURY):
        r = 2 + k
        rows.append(row_xml(r, [cell("%s%d" % (col_letters(j + 1), r), v, ST_LICZBA if j == 4 else ST_TEKST) for j, v in enumerate(rec)]))
    info = 2 + len(RECEPTURY) + 2
    rows.append(row_xml(info, [cell("A%d" % info, "PRODUKCJA zużywa SUROWIEC z tej tabeli (np. 120 MP zrębki ÷ 4 = 30 m³ drewna). Zrębka nie może być surowcem. Jednostki: MP, M3, TON.", ST_TEKST)]))
    nowe_ark.append(("RECEPTURY", "sheet7.xml", "Arkusz7",
                     worksheet("Arkusz7", [(1, 34), (2, 12), (3, 36), (4, 12), (5, 22), (6, 14), (7, 40)], "".join(rows), tab_color="FF2F75B5"), False))

    # HISTORIA_ZMIAN
    nag = ["Data i godzina", "Użytkownik", "ID operacji", "Akcja", "Pole", "Było", "Jest", "Powód"]
    tak = sum(1 for m in migr if m["Transport"] == "TAK")
    zlozone = len({m["ID_operacji"] for m in migr if m["Rola_wiersza"] != GLOWNY})
    wpisy = [
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "Kolumny Tabela_dane", "A:AL (38)", "A:AT (46): " + ", ".join(NOWE_KOLUMNY), "Transport jako część operacji, receptury produkcji"),
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "Identyfikatory operacji", "brak", "%d wierszy → %d operacji H-xxxx (w tym %d złożonych: wiersz główny + wiersze automatyczne)" % (len(migr), len({m['ID_operacji'] for m in migr}), zlozone), ""),
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "Transport", "typ operacji TRANSPORT / kolumny R-U", "pole Transport: TAK w %d wierszach, NIE w %d" % (tak, len(migr) - tak), "dane historyczne bez zmian"),
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "Receptury produkcji", "zużycie = ta sama ilość i produkt „Zrzyna”/zrębka", "arkusz RECEPTURY (%d receptur, 1 m³ = 4 MP)" % len(RECEPTURY), "zużycie drewna zamiast zrębki"),
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "MAGAZYN!AA12:AC12", 'COUNTIF/SUMIF(typ="TRANSPORT")', "COUNTIFS/SUMIFS(Transport=TAK)", ""),
        (teraz, "System (migracja 4.0)", "—", "MIGRACJA", "SŁOWNIK!F7", "TRANSPORT", "(usunięto z listy typów)", ""),
    ]
    rows = [row_xml(1, [cell("%s1" % col_letters(i + 1), h, ST_NAGLOWEK) for i, h in enumerate(nag)], 30)]
    for k, w in enumerate(wpisy):
        r = 2 + k
        rows.append(row_xml(r, [cell("%s%d" % (col_letters(j + 1), r), excel_serial(v) if j == 0 else v, ST_CZAS if j == 0 else ST_TEKST) for j, v in enumerate(w)]))
    nowe_ark.append(("HISTORIA_ZMIAN", "sheet8.xml", "Arkusz8",
                     worksheet("Arkusz8", [(1, 18), (2, 22), (3, 12), (4, 14), (5, 28), (6, 40), (7, 40), (8, 34)], "".join(rows), tab_color="FF7F6000", freeze=2), False))

    # ARCHIWUM
    naglowki_dane = [c.value for c in ws[1]][:38] + NOWE_KOLUMNY + ["Data archiwizacji", "Użytkownik", "Akcja", "Powód"]
    rows = [row_xml(1, [cell("%s1" % col_letters(i + 1), h, ST_NAGLOWEK) for i, h in enumerate(naglowki_dane)], 30)]
    nowe_ark.append(("ARCHIWUM", "sheet9.xml", "Arkusz9",
                     worksheet("Arkusz9", [(i, 16) for i in range(1, 51)], "".join(rows), tab_color="FF808080", freeze=2), False))

    # DOKUMENT
    rows = [row_xml(1, [cell("A1", "DOKUMENT OPERACJI MAGAZYNOWEJ", ST_NAGLOWEK)], 26),
            row_xml(3, [cell("A3", "Wybierz operację w arkuszu REJESTR_OPERACJI i kliknij „DOKUMENT PDF”.", ST_TEKST)])]
    nowe_ark.append(("DOKUMENT", "sheet10.xml", "Arkusz10",
                     worksheet("Arkusz10", [(1, 26), (2, 40), (3, 12), (4, 8), (5, 12), (6, 14), (7, 14)], "".join(rows)), False))

    # przyciski na REJESTR_OPERACJI (kształty z makrami)
    przyciski = [("NOWA OPERACJA", "OtworzFormularz", "1E6B45"), ("KOREKTA", "KorektaOperacji", "2F75B5"),
                 ("USUŃ", "UsunOperacjeUI", "C00000"), ("DOKUMENT PDF", "DokumentOperacjiPDF", "7F6000"), ("ODŚWIEŻ", "OdswiezRejestrUI", "595959")]
    sp = ""
    for i, (tekst, makro, kolor) in enumerate(przyciski):
        c0 = 2 + i * 2
        sp += ('<xdr:twoCellAnchor editAs="absolute"><xdr:from><xdr:col>%d</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>30000</xdr:rowOff></xdr:from>'
               '<xdr:to><xdr:col>%d</xdr:col><xdr:colOff>600000</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
               '<xdr:sp macro="[0]!%s" textlink=""><xdr:nvSpPr><xdr:cNvPr id="%d" name="Przycisk %s"/><xdr:cNvSpPr/></xdr:nvSpPr>'
               '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
               '<a:solidFill><a:srgbClr val="%s"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
               '<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
               '<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="pl-PL" sz="1100" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>%s</a:t></a:r></a:p></xdr:txBody></xdr:sp>'
               '<xdr:clientData/></xdr:twoCellAnchor>') % (c0, c0 + 1, makro, i + 2, esc(tekst), kolor, esc(tekst))
    files["xl/drawings/drawing4.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">%s</xdr:wsDr>' % sp).encode("utf-8")

    ct = files["[Content_Types].xml"].decode("utf-8")
    wbx = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    sheet_id = max(int(x) for x in re.findall(r'sheetId="(\d+)"', wbx))
    rel_ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
    nowe_sheets = ""
    for nazwa, plik, code, xml, ma_rys in nowe_ark:
        files["xl/worksheets/" + plik] = xml.encode("utf-8")
        if ma_rys:
            files["xl/worksheets/_rels/%s.rels" % plik] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing4.xml"/>'
                '</Relationships>').encode("utf-8")
        ct = ct.replace("</Types>", '<Override PartName="/xl/worksheets/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>' % plik)
        sheet_id += 1
        rid = max(rel_ids) + 1
        rel_ids.append(rid)
        rels = rels.replace("</Relationships>", '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/%s"/></Relationships>' % (rid, plik))
        nowe_sheets += '<sheet name="%s" sheetId="%d" r:id="rId%d"/>' % (nazwa, sheet_id, rid)
    ct = ct.replace("</Types>", '<Override PartName="/xl/drawings/drawing4.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    # nowe karty na końcu (indeksy istniejących arkuszy bez zmian - tabele przestawne, fragmentatory, nazwy)
    wbx = wbx.replace('<sheet name="SŁOWNIK" sheetId="11" r:id="rId5"/>', '<sheet name="SŁOWNIK" sheetId="11" r:id="rId5"/>' + nowe_sheets, 1)
    assert wbx.count("<sheet ") == 10, wbx.count("<sheet ")
    wbx = wbx.replace("Dane!$A$1:$AL$%d" % ostatni, "Dane!$A$1:$AT$%d" % ostatni)
    wbx = wbx.replace("<definedName name=\"DV_E\">'SŁOWNIK'!$F$2:$F$7", "<definedName name=\"DV_E\">'SŁOWNIK'!$F$2:$F$6")
    # tabele przestawne / aktywna karta: pozostaje Dane
    files["xl/workbook.xml"] = wbx.encode("utf-8")
    files["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    # calcChain - Excel odbuduje
    files.pop("xl/calcChain.xml", None)
    ct = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", ct)
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rels = re.sub(r'<Relationship [^>]*Target="calcChain.xml"/>', "", rels)
    files["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    files["[Content_Types].xml"] = ct.encode("utf-8")

    app = files["docProps/app.xml"].decode("utf-8")
    nazwy = ["Dane", "MAGAZYN", "KONTRAHENCI", "TRANSPORT", "SŁOWNIK", "REJESTR_OPERACJI", "RECEPTURY", "HISTORIA_ZMIAN", "ARCHIWUM", "DOKUMENT"]
    app = re.sub(r"<vt:i4>5</vt:i4>", "<vt:i4>10</vt:i4>", app)
    app = re.sub(r'<TitlesOfParts>.*?</TitlesOfParts>', '<TitlesOfParts><vt:vector size="10" baseType="lpstr">%s</vt:vector></TitlesOfParts>'
                 % "".join("<vt:lpstr>%s</vt:lpstr>" % n for n in nazwy), app)
    files["docProps/app.xml"] = app.encode("utf-8")

    # ---------------- 5. projekt VBA ----------------
    tmp_bin = out_path + ".vbaProject.bin"
    with open(tmp_bin, "wb") as f:
        f.write(files["xl/vbaProject.bin"])
    orig_prj = vbaproj.read_project(tmp_bin)
    zrodlo = {m["name"]: m for m in orig_prj["modules"]}

    def plik(n):
        return open(os.path.join(src_dir, n), encoding="utf-8").read()

    def naglowek(m):
        return "\r\n".join(l for l in m["source"].replace("\r\n", "\n").split("\n") if l.startswith("Attribute ")) + "\r\n"

    def dokument(nazwa, kod=""):
        return ('Attribute VB_Name = "%s"\r\nAttribute VB_Base = "0{00020820-0000-0000-C000-000000000046}"\r\n'
                'Attribute VB_GlobalNameSpace = False\r\nAttribute VB_Creatable = False\r\nAttribute VB_PredeclaredId = True\r\n'
                'Attribute VB_Exposed = True\r\nAttribute VB_TemplateDerived = False\r\nAttribute VB_Customizable = True\r\n%s' % (nazwa, kod))

    moduly = []
    for m in orig_prj["modules"]:
        n = m["name"]
        if n == "frmDane":
            moduly.append({"name": n, "kind": "form", "source": naglowek(m) + plik("frmDane.frm")})
        elif n == "frmDodatkowe":
            moduly.append({"name": n, "kind": "form", "source": naglowek(m) + plik("frmDodatkowe.frm")})
        elif n == "Module1":
            moduly.append({"name": n, "kind": "std", "source": 'Attribute VB_Name = "Module1"\r\n' + plik("Module1.bas")})
        else:
            moduly.append({"name": n, "kind": "doc", "source": m["source"]})
    moduly.append({"name": "Arkusz6", "kind": "doc", "source": dokument("Arkusz6", plik("Arkusz6_REJESTR.cls"))})
    for k in ("Arkusz7", "Arkusz8", "Arkusz9", "Arkusz10"):
        moduly.append({"name": k, "kind": "doc", "source": dokument(k)})
    moduly.append({"name": "clsOperacja", "kind": "class", "source": 'Attribute VB_Name = "clsOperacja"\r\nAttribute VB_GlobalNameSpace = False\r\n'
                   'Attribute VB_Creatable = False\r\nAttribute VB_PredeclaredId = False\r\nAttribute VB_Exposed = False\r\n' + plik("clsOperacja.cls")})
    moduly.append({"name": "modMagazyn", "kind": "std", "source": 'Attribute VB_Name = "modMagazyn"\r\n' + plik("modMagazyn.bas")})
    moduly.append({"name": "modTesty", "kind": "std", "source": 'Attribute VB_Name = "modTesty"\r\n' + plik("modTesty.bas")})
    vbaproj.build_project(orig_prj, moduly, tmp_bin)
    files["xl/vbaProject.bin"] = open(tmp_bin, "rb").read()
    os.remove(tmp_bin)

    # ---------------- zapis ----------------
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            zi = infos.get(name)
            if zi is None:
                zi = zipfile.ZipInfo(name, date_time=teraz.timetuple()[:6])
                zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    return {"wiersze": len(migr), "operacje": len({m["ID_operacji"] for m in migr}), "transport_tak": tak, "rejestr": len(rej)}


if __name__ == "__main__":
    print(build(sys.argv[1], sys.argv[2], sys.argv[3]))
